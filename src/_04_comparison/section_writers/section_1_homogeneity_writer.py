"""
Section 1 Writer - Homogeneity Analysis

Creates Section 1: Homogenität (4 sheets) showing:
- 1a: Cluster statistics tables per algorithm
- 1b: Cluster size charts and score distributions
- 1c: Homogeneity metrics by sector
- 1d: Outlier analysis

Addresses research question: Bilden kennzahlenbasierte Cluster intern homogenere Gruppen?
"""

import pandas as pd
import logging
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from .base_section_writer import BaseSectionWriter

logger = logging.getLogger(__name__)


class Section1HomogeneityWriter(BaseSectionWriter):
    """Writer for Section 1: Homogeneity Analysis"""

    def __init__(self, config, colors=None, score_columns=None, market='germany'):
        """
        Initialize homogeneity writer

        Args:
            config: Configuration dictionary
            colors: Color scheme dictionary
            score_columns: List of score column names
            market: Market name for PNG paths
        """
        super().__init__(config, colors, score_columns)
        self.market = market

    def create_section(self, wb: Workbook, df: pd.DataFrame):
        """
        Create Section 1: Homogeneity (4 sheets)

        Args:
            wb: Workbook object
            df: Overview dataframe with all algorithm results
        """
        self._create_section_1a_tables(wb, df)
        self._create_section_1b_charts(wb, df)
        self._create_section_1c_sector(wb, df)
        self._create_section_1d_outliers(wb, df)

    def _create_section_1a_tables(self, wb: Workbook, df: pd.DataFrame):
        """Section 1a: Homogenität Tables"""
        ws = wb.create_sheet("1a_Homogenität_Tabellen")

        # Title
        ws['A1'] = "SECTION 1a: HOMOGENITÄT - TABLES"
        ws['A1'].font = Font(size=14, bold=True, color="FFFFFF")
        ws['A1'].fill = PatternFill(start_color=self.colors['header'], fill_type='solid')
        ws.merge_cells('A1:J1')

        row = 3

        # Interpretation box - Kernfrage 1
        row = self._add_interpretation_box(ws, row, "KERNFRAGE 1: HOMOGENITÄT", [
            "Bilden kennzahlenbasierte Cluster intern homogenere Gruppen als traditionelle Klassifikationen?",
            "Silhouette Score > 0.5 zeigt gute interne Homogenität (Werte zwischen -1 und 1)",
            "Niedrige Intra-Cluster-Varianz = Unternehmen im Cluster sind sich ähnlich",
            "Vergleich: Cluster-Homogenität vs. GICS-Branchen-Homogenität (siehe Section 2)"
        ], merge_cols=10)
        row += 1

        # ===== CLUSTER STATISTICS PER ALGORITHM =====
        for algo in df['algorithm'].unique():
            algo_df = df[df['algorithm'] == algo].copy()

            ws[f'A{row}'] = f"{algo.upper()} - CLUSTER STATISTICS"
            ws[f'A{row}'].font = Font(size=12, bold=True)
            ws[f'A{row}'].fill = PatternFill(start_color=self.colors[algo], fill_type='solid')
            ws.merge_cells(f'A{row}:J{row}')
            row += 1

            # Group by cluster
            cluster_stats = []
            for cluster_id in sorted(algo_df['cluster'].unique()):
                cluster_df = algo_df[algo_df['cluster'] == cluster_id]

                stats = {
                    'Cluster ID': cluster_id,
                    'Size (N)': len(cluster_df),
                    'Size (%)': f"{len(cluster_df)/len(algo_df)*100:.1f}%"
                }

                # Add score statistics with improved headers
                score_name_mapping = {
                    'overall_score': 'Overall Score',
                    'proximity_score': 'Proximity Score',
                    'profitability_score': 'Profitability Score',
                    'leverage_score': 'Leverage Score',
                    'efficiency_score': 'Efficiency Score',
                    'growth_score': 'Growth Score',
                    'relative_score': 'Relative Score'
                }

                for score_col in self.score_columns:
                    if score_col in cluster_df.columns:
                        display_name = score_name_mapping.get(score_col, score_col)
                        stats[f'{display_name} (Ø)'] = cluster_df[score_col].mean()
                        stats[f'{display_name} (σ)'] = cluster_df[score_col].std()

                cluster_stats.append(stats)

            # Convert to dataframe
            stats_df = pd.DataFrame(cluster_stats)

            # Write to sheet
            for r_idx, row_data in enumerate(dataframe_to_rows(stats_df, index=False, header=True)):
                for c_idx, value in enumerate(row_data):
                    cell = ws.cell(row=row, column=c_idx+1, value=value)
                    if r_idx == 0:  # Header
                        cell.font = Font(bold=True)
                        cell.fill = PatternFill(start_color=self.colors['neutral'], fill_type='solid')
                    elif isinstance(value, (int, float)) and not isinstance(value, bool):
                        cell.number_format = '0.00'
                row += 1

            row += 2

        # Column widths
        for col in range(1, 11):
            ws.column_dimensions[chr(64+col)].width = 15

        # ===== EMBED VISUALIZATIONS =====
        row += 2
        ws[f'A{row}'] = "CLUSTER QUALITY VISUALIZATIONS"
        ws[f'A{row}'].font = Font(size=12, bold=True)
        ws.merge_cells(f'A{row}:J{row}')
        row += 1

        # Try to embed cluster quality plots from K-Means combined
        base_path = Path(f'output/{self.market}/02_algorithms/kmeans_comparative/combined')
        plots_to_embed = [
            ('1_cluster_quality/plots/cluster_homogeneity.png', 'A', 0.4),
            ('1_cluster_quality/plots/score_distribution_overall.png', 'F', 0.4),
            ('4_company_insights/plots/cluster_sizes.png', 'K', 0.4),
        ]

        embedded_count = 0
        for plot_path, col_letter, scale in plots_to_embed:
            full_path = base_path / plot_path
            if full_path.exists():
                self._embed_png(ws, full_path, f'{col_letter}{row}', scale=scale)
                embedded_count += 1

        if embedded_count > 0:
            row += 30  # Space for embedded images
        else:
            ws[f'A{row}'] = "Cluster quality visualizations not available"
            ws[f'A{row}'].font = Font(italic=True, size=9)
            row += 1

        logger.info(f"  ✓ Section 1a sheet created ({embedded_count} plots embedded)")

    def _create_section_1b_charts(self, wb: Workbook, df: pd.DataFrame):
        """Section 1b: Homogenität Charts"""
        ws = wb.create_sheet("1b_Homogenität_Charts")

        # Title
        ws['A1'] = "SECTION 1b: HOMOGENITÄT - CHARTS"
        ws['A1'].font = Font(size=14, bold=True, color="FFFFFF")
        ws['A1'].fill = PatternFill(start_color=self.colors['header'], fill_type='solid')
        ws.merge_cells('A1:H1')

        row = 3

        # ===== CLUSTER SIZE BAR CHART =====
        ws[f'A{row}'] = "CLUSTER SIZES BY ALGORITHM"
        ws[f'A{row}'].font = Font(size=12, bold=True)
        row += 1

        # Prepare data
        cluster_sizes = df.groupby(['algorithm', 'cluster']).size().unstack(fill_value=0)

        # Write data
        start_row = row
        for r_idx, row_data in enumerate(dataframe_to_rows(cluster_sizes, index=True, header=True)):
            for c_idx, value in enumerate(row_data):
                cell = ws.cell(row=row, column=c_idx+1, value=value)
                if r_idx == 0:  # Header
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(start_color=self.colors['neutral'], fill_type='solid')
            row += 1

        # Add bar chart
        chart = BarChart()
        chart.title = "Cluster Sizes"
        chart.x_axis.title = "Algorithm"
        chart.y_axis.title = "Number of Companies"
        chart.type = "col"
        chart.grouping = "clustered"

        data = Reference(ws, min_col=2, max_col=cluster_sizes.shape[1]+1,
                        min_row=start_row, max_row=row-1)
        cats = Reference(ws, min_col=1, min_row=start_row+1, max_row=row-1)

        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height = 15
        chart.width = 25

        ws.add_chart(chart, 'A' + str(row + 2))

        row += 20

        # ===== EMBED PERFORMANCE DASHBOARD & VISUALIZATIONS =====
        ws[f'A{row}'] = "PERFORMANCE & SCORE VISUALIZATIONS"
        ws[f'A{row}'].font = Font(size=12, bold=True)
        ws.merge_cells(f'A{row}:H{row}')
        row += 1

        # Embed from K-Means combined (main analysis)
        base_path = Path(f'output/{self.market}/02_algorithms/kmeans_comparative/combined')
        plots_to_embed = [
            ('plots/performance_dashboard.png', 'A', 0.35),
            ('1_cluster_quality/plots/score_correlations.png', 'J', 0.35),
            ('plots/cluster_distribution.png', 'S', 0.35),
        ]

        embedded_count = 0
        for plot_path, col_letter, scale in plots_to_embed:
            full_path = base_path / plot_path
            if full_path.exists():
                self._embed_png(ws, full_path, f'{col_letter}{row}', scale=scale)
                embedded_count += 1

        if embedded_count == 0:
            ws[f'A{row}'] = "Performance visualizations not available"
            ws[f'A{row}'].font = Font(italic=True, size=9)
            row += 1
        else:
            row += 28  # Space for embedded images

        # ===== SCORE DISTRIBUTION TABLE =====
        if 'overall_score' in df.columns:
            ws[f'A{row}'] = "OVERALL SCORE DISTRIBUTION"
            ws[f'A{row}'].font = Font(size=12, bold=True)
            row += 1

            # Calculate percentiles
            score_dist = df.groupby('algorithm')['overall_score'].describe()

            # Write data
            for r_idx, row_data in enumerate(dataframe_to_rows(score_dist, index=True, header=True)):
                for c_idx, value in enumerate(row_data):
                    cell = ws.cell(row=row, column=c_idx+1, value=value)
                    if r_idx == 0:
                        cell.font = Font(bold=True)
                        cell.fill = PatternFill(start_color=self.colors['neutral'], fill_type='solid')
                    elif isinstance(value, float):
                        cell.number_format = '0.00'
                row += 1

        logger.info("  ✓ Section 1b sheet created")

    def _create_section_1c_sector(self, wb: Workbook, df: pd.DataFrame):
        """Section 1c: Homogenität by Sector"""
        ws = wb.create_sheet("1c_Homogenität_Sektor")

        # Title
        ws['A1'] = "SECTION 1c: HOMOGENITÄT BY SECTOR"
        ws['A1'].font = Font(size=14, bold=True, color="FFFFFF")
        ws['A1'].fill = PatternFill(start_color=self.colors['header'], fill_type='solid')
        ws.merge_cells('A1:F1')

        if 'gsector' not in df.columns:
            ws['A3'] = "No sector data available"
            logger.warning("  ⚠ No sector data for Section 1c")
            return

        row = 3

        # ===== CLUSTER QUALITY BY SECTOR =====
        ws[f'A{row}'] = "CLUSTER QUALITY METRICS BY SECTOR"
        ws[f'A{row}'].font = Font(size=12, bold=True)
        ws.merge_cells(f'A{row}:F{row}')
        row += 1

        # Calculate metrics per sector
        sector_metrics = []
        for sector in sorted(df['gsector'].unique()):
            if pd.isna(sector):
                continue

            sector_df = df[df['gsector'] == sector]

            metrics = {
                'GICS Sector': sector,
                'Companies (N)': len(sector_df),
                'Clusters (N)': sector_df['cluster'].nunique(),
                'Outlier Rate (%)': f"{sector_df['is_outlier'].mean()*100:.1f}" if 'is_outlier' in sector_df.columns else 'N/A'
            }

            # Add score statistics
            if 'overall_score' in sector_df.columns:
                metrics['Avg Score (Ø)'] = sector_df['overall_score'].mean()
                metrics['Score StdDev (σ)'] = sector_df['overall_score'].std()

            sector_metrics.append(metrics)

        # Convert to dataframe
        sector_df = pd.DataFrame(sector_metrics)

        # Write to sheet
        start_row = row
        for r_idx, row_data in enumerate(dataframe_to_rows(sector_df, index=False, header=True)):
            for c_idx, value in enumerate(row_data):
                cell = ws.cell(row=row, column=c_idx+1, value=value)
                if r_idx == 0:  # Header
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(start_color=self.colors['neutral'], fill_type='solid')
                elif isinstance(value, float):
                    cell.number_format = '0.00'
            row += 1

        # Add conditional formatting for scores
        sector_df_check = pd.DataFrame(sector_metrics)
        if 'Avg Score (Ø)' in sector_df_check.columns:
            score_col = chr(64 + sector_df_check.columns.get_loc('Avg Score (Ø)') + 1)
            ws.conditional_formatting.add(
                f'{score_col}{start_row+1}:{score_col}{row-1}',
                ColorScaleRule(
                    start_type='min', start_color='FFC7CE',
                    mid_type='percentile', mid_value=50, mid_color='FFEB9C',
                    end_type='max', end_color='C6EFCE'
                )
            )

            # Add color legend
            row += 1
            ws[f'A{row}'] = "Color Legend: 🔴 Red = Low Score | 🟡 Yellow = Medium | 🟢 Green = High Score"
            ws[f'A{row}'].font = Font(size=9, italic=True, color='666666')
            ws.merge_cells(f'A{row}:F{row}')

        # Column widths
        for col in range(1, 7):
            ws.column_dimensions[chr(64+col)].width = 18

        # ===== EMBED CLUSTER PROFILE VISUALIZATIONS =====
        row += 2
        ws[f'A{row}'] = "CLUSTER PROFILE VISUALIZATIONS"
        ws[f'A{row}'].font = Font(size=12, bold=True)
        ws.merge_cells(f'A{row}:F{row}')
        row += 1

        # Embed cluster characteristics and correlation heatmap
        base_path = Path(f'output/{self.market}/02_algorithms/kmeans_comparative/combined')
        plots_to_embed = [
            ('plots/cluster_characteristics.png', 'A', 0.4),
            ('plots/correlation_heatmap.png', 'G', 0.35),
        ]

        embedded_count = 0
        for plot_path, col_letter, scale in plots_to_embed:
            full_path = base_path / plot_path
            if full_path.exists():
                self._embed_png(ws, full_path, f'{col_letter}{row}', scale=scale)
                embedded_count += 1

        if embedded_count > 0:
            row += 30  # Space for embedded images

        logger.info(f"  ✓ Section 1c sheet created ({embedded_count} plots embedded)")

    def _create_section_1d_outliers(self, wb: Workbook, df: pd.DataFrame):
        """Section 1d: Outlier Analysis"""
        ws = wb.create_sheet("1d_Homogenität_Outliers")

        # Title
        ws['A1'] = "SECTION 1d: OUTLIER ANALYSIS"
        ws['A1'].font = Font(size=14, bold=True, color="FFFFFF")
        ws['A1'].fill = PatternFill(start_color=self.colors['header'], fill_type='solid')
        ws.merge_cells('A1:H1')

        if 'is_outlier' not in df.columns:
            ws['A3'] = "No outlier data available"
            logger.warning("  ⚠ No outlier data for Section 1d")
            return

        row = 3

        # ===== OUTLIER SUMMARY =====
        ws[f'A{row}'] = "OUTLIER SUMMARY"
        ws[f'A{row}'].font = Font(size=12, bold=True)
        ws[f'A{row}'].fill = PatternFill(start_color=self.colors['subheader'], fill_type='solid')
        ws.merge_cells(f'A{row}:B{row}')
        row += 1

        outliers_df = df[df['is_outlier'] == True]

        summary = [
            ("Total Outliers", len(outliers_df)),
            ("Outlier Rate (%)", f"{len(outliers_df)/len(df)*100:.1f}%"),
            ("DBSCAN Noise", len(outliers_df[outliers_df['outlier_reason'].str.contains('Noise', na=False)])),
            ("Low Score", len(outliers_df[outliers_df['outlier_reason'].str.contains('Low Score', na=False)])),
            ("High Variance", len(outliers_df[outliers_df['outlier_reason'].str.contains('High Variance', na=False)])),
        ]

        for label, value in summary:
            ws[f'A{row}'] = label
            ws[f'B{row}'] = value
            ws[f'A{row}'].font = Font(bold=True)
            row += 1

        row += 2

        # ===== OUTLIER DETAILS =====
        ws[f'A{row}'] = "OUTLIER COMPANIES"
        ws[f'A{row}'].font = Font(size=12, bold=True)
        ws.merge_cells(f'A{row}:H{row}')
        row += 1

        # Select relevant columns
        outlier_cols = ['conm', 'gvkey', 'algorithm', 'cluster', 'outlier_reason']
        if 'overall_score' in outliers_df.columns:
            outlier_cols.append('overall_score')
        if 'gsector' in outliers_df.columns:
            outlier_cols.append('gsector')

        outlier_display = outliers_df[outlier_cols].copy()

        # Write to sheet
        for r_idx, row_data in enumerate(dataframe_to_rows(outlier_display, index=False, header=True)):
            for c_idx, value in enumerate(row_data):
                cell = ws.cell(row=row, column=c_idx+1, value=value)
                if r_idx == 0:  # Header
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(start_color=self.colors['bad'], fill_type='solid')
                elif isinstance(value, float):
                    cell.number_format = '0.00'
            row += 1

        # Column widths
        ws.column_dimensions['A'].width = 30
        for col in range(2, 9):
            ws.column_dimensions[chr(64+col)].width = 15

        # ===== EMBED OUTLIER & PERFORMANCE VISUALIZATIONS =====
        row += 2
        ws[f'A{row}'] = "OUTLIER & PERFORMANCE VISUALIZATIONS"
        ws[f'A{row}'].font = Font(size=12, bold=True)
        ws.merge_cells(f'A{row}:H{row}')
        row += 1

        # Embed outliers and top/bottom performers plots
        base_path = Path(f'output/{self.market}/02_algorithms/kmeans_comparative/combined')
        plots_to_embed = [
            ('4_company_insights/plots/outliers.png', 'A', 0.4),
            ('4_company_insights/plots/top_bottom_performers.png', 'F', 0.4),
        ]

        embedded_count = 0
        for plot_path, col_letter, scale in plots_to_embed:
            full_path = base_path / plot_path
            if full_path.exists():
                self._embed_png(ws, full_path, f'{col_letter}{row}', scale=scale)
                embedded_count += 1

        if embedded_count > 0:
            row += 30  # Space for embedded images

        logger.info(f"  ✓ Section 1d sheet created ({embedded_count} plots embedded)")
