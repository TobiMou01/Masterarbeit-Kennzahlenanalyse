"""
Base Section Writer - Shared utilities for all section writers

Provides common functionality for creating Excel sections:
- Styling (colors, fonts, borders)
- Helper methods (embedding images, interpretation boxes)
- Shared constants and configurations
"""

import logging
from pathlib import Path
from typing import Dict, List
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from PIL import Image
import io

logger = logging.getLogger(__name__)


class BaseSectionWriter:
    """Base class for all section writers with shared utilities"""

    def __init__(self, config: Dict, colors: Dict = None, score_columns: List[str] = None):
        """
        Initialize base section writer

        Args:
            config: Configuration dictionary
            colors: Color scheme dictionary (optional)
            score_columns: List of score column names (optional)
        """
        self.config = config

        # Default color scheme
        self.colors = colors or {
            'header': 'FF1F4E78',      # Dark blue header
            'subheader': 'FF4472C4',   # Medium blue
            'config': 'FFF4B084',      # Orange for config boxes
            'good': 'FFC6EFCE',        # Light green
            'warning': 'FFFFF2CC',     # Light yellow
            'bad': 'FFFFC7CE',         # Light red
            'neutral': 'FFE7E6E6',     # Light gray
            'kmeans': 'FFCCE5FF',      # Light blue
            'hierarchical': 'FFCCFFCC', # Light green
            'dbscan': 'FFFFCCCC'       # Light red
        }

        # Default score columns
        self.score_columns = score_columns or [
            'proximity_score',
            'profitability_score',
            'leverage_score',
            'efficiency_score',
            'growth_score',
            'relative_score',
            'overall_score'
        ]

    def _embed_png(self, ws, png_path: Path, cell: str, scale: float = 0.5):
        """
        Embed PNG image into worksheet

        Args:
            ws: Worksheet object
            png_path: Path to PNG file
            cell: Cell reference (e.g., 'A1')
            scale: Scale factor for image (default 0.5)

        Returns:
            True if successful, False otherwise
        """
        try:
            if not png_path.exists():
                logger.warning(f"PNG not found: {png_path}")
                return False

            # Load and resize image
            img = Image.open(png_path)

            # Resize if needed
            if scale != 1.0:
                new_size = (int(img.width * scale), int(img.height * scale))
                img = img.resize(new_size, Image.Resampling.LANCZOS)

            # Save to BytesIO
            img_byte_arr = io.BytesIO()
            img.save(img_byte_arr, format='PNG')
            img_byte_arr.seek(0)

            # Create Excel image
            xl_img = XLImage(img_byte_arr)
            ws.add_image(xl_img, cell)

            logger.info(f"  ✓ PNG embedded: {png_path.name}")
            return True

        except Exception as e:
            logger.error(f"Error embedding PNG {png_path}: {e}")
            return False

    def _add_interpretation_box(self, ws, row: int, title: str, points: List[str], merge_cols: int = 6) -> int:
        """
        Add interpretation help box to worksheet

        Args:
            ws: Worksheet object
            row: Starting row number
            title: Box title (e.g., "KERNFRAGE 1: HOMOGENITÄT")
            points: List of interpretation points
            merge_cols: Number of columns to merge (default 6)

        Returns:
            Next row number after the box
        """
        # Title
        ws[f'A{row}'] = f"📊 {title}"
        ws[f'A{row}'].font = Font(size=11, bold=True, color="1F4E78")
        ws[f'A{row}'].fill = PatternFill(start_color='FFF4E6', fill_type='solid')
        ws[f'A{row}'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        ws.merge_cells(f'A{row}:{chr(64+merge_cols)}{row}')

        # Add border
        border = Border(
            left=Side(style='medium', color='1F4E78'),
            right=Side(style='medium', color='1F4E78'),
            top=Side(style='medium', color='1F4E78'),
            bottom=Side(style='thin', color='1F4E78')
        )
        ws[f'A{row}'].border = border
        row += 1

        # Interpretation points
        for point in points:
            ws[f'A{row}'] = f"  • {point}"
            ws[f'A{row}'].font = Font(size=10)
            ws[f'A{row}'].fill = PatternFill(start_color='FFF9F0', fill_type='solid')
            ws[f'A{row}'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            ws.merge_cells(f'A{row}:{chr(64+merge_cols)}{row}')

            # Border for content
            border = Border(
                left=Side(style='medium', color='1F4E78'),
                right=Side(style='medium', color='1F4E78'),
                bottom=Side(style='thin', color='E7E6E6')
            )
            ws[f'A{row}'].border = border
            row += 1

        # Bottom border for last row
        last_row = row - 1
        ws[f'A{last_row}'].border = Border(
            left=Side(style='medium', color='1F4E78'),
            right=Side(style='medium', color='1F4E78'),
            bottom=Side(style='medium', color='1F4E78')
        )

        # Set row heights for better readability
        for r in range(row - len(points) - 1, row):
            ws.row_dimensions[r].height = 25

        return row + 1

    def _apply_border(self, ws, cell_range: str, style: str = 'thin'):
        """
        Apply border to cell range

        Args:
            ws: Worksheet object
            cell_range: Cell range (e.g., 'A1:D10')
            style: Border style ('thin', 'medium', 'thick')
        """
        border = Border(
            left=Side(style=style),
            right=Side(style=style),
            top=Side(style=style),
            bottom=Side(style=style)
        )

        for row in ws[cell_range]:
            for cell in row:
                cell.border = border
